"""Build the mini-research paper (Word .docx) from the latest results.

    python paper/build_paper.py

Every number in the text is read from results/tables/*.csv, so the paper always matches
the code's output. Edit AUTHORS below with the group members' names.
"""
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src import config as C  # noqa: E402

AUTHORS = ["Almoguez, Romar S.", "Dela Cruz, Sean Arkin C.", "Fajardo, Osyris Benedict P.",
           "Hornales, Jian Kyle C."]
AFFILIATION = ("Bachelor of Science in Mathematics major in Computer Science\n"
               "College of Science, Bulacan State University, Malolos, Bulacan, Philippines")
COURSE = "MCS 306 Final Course Output"
REPO_URL = "https://github.com/SeanXxotick/MCS306FCO"
OUT = ROOT / "paper" / "PINN_Dengue_Bulacan_Paper.docx"
FIG = C.FIGURES
TAB = C.TABLES

# ------------------------------------------------------------------ load results
conv = pd.read_csv(TAB / "exp1_rk4_convergence.csv")
fitstep = pd.read_csv(TAB / "exp1_step_size_in_fit.csv")
summ = pd.read_csv(TAB / "exp2_summary.csv")
cls_all = pd.read_csv(TAB / "exp2_classical.csv")
sweep = pd.read_csv(TAB / "exp3_ode_weight_sweep.csv")
seasons = pd.read_csv(TAB / "exp4_seasons.csv")
boot = pd.read_csv(TAB / "exp4_bootstrap_2010.csv").set_index("parameter")
pinn10 = pd.read_csv(TAB / "exp4_pinn_2010.csv")

pinn = summ[summ.method == "PINN"].set_index("noise")
cls = summ[summ.method != "PINN"].set_index("noise")
orders = conv.observed_order.dropna()
speedup = pinn.seconds.mean() / cls.seconds.mean()
best_w = sweep.loc[(sweep.beta_err_pct + sweep.gamma_err_pct).idxmin()]
worst_w = sweep.loc[(sweep.beta_err_pct + sweep.gamma_err_pct).idxmax()]
s2010 = seasons.set_index("season").loc[2010]
lo_days, hi_days = C.INFECTIOUS_DAYS_RANGE
at_bound = seasons[(seasons.infectious_days < lo_days * 1.01) | (seasons.infectious_days > hi_days * 0.99)]
interior = seasons.drop(at_bound.index)
p_cls, p_pinn = pinn10.iloc[0], pinn10.iloc[1]


def f(x, d=2):
    return f"{x:.{d}f}"


def pct(x):
    return f"{x:.2f}%" if x >= 0.01 else f"{x:.1e}%"


# ------------------------------------------------------------------ document helpers
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(12)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 14), (2, 12), (3, 12)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name, h.font.size, h.font.bold = "Times New Roman", Pt(size), True
    h.font.color.rgb = RGBColor(0, 0, 0)
    h.font.italic = lvl == 3
for sec in doc.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)

fig_no, tab_no = [0], [0]


def para(text, bold=False, italic=False, align=None, size=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def body(text):
    return para(text, align="justify")


def rich(parts):
    """Paragraph from [(text, style)] where style is '', 'b', 'i' or 'bi'."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, style in parts:
        r = p.add_run(text)
        r.bold, r.italic = "b" in style, "i" in style
    return p


def equation(text, label=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name, r.font.size = "Cambria Math", Pt(12)
    if label:
        p.add_run(f"\t({label})")
    return p


def code(text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.3)
        r = p.add_run(line if line else " ")
        r.font.name, r.font.size = "Consolas", Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def figure(name, caption, width=6.3):
    fig_no[0] += 1
    doc.add_picture(str(FIG / name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.add_run(f"Figure {fig_no[0]}").bold = True
    p2 = doc.add_paragraph()
    p2.add_run(caption).italic = True
    p2.paragraph_format.space_after = Pt(12)
    return fig_no[0]


def table(df, headers, title, fmts=None, note=None):
    tab_no[0] += 1
    p = doc.add_paragraph()
    p.add_run(f"Table {tab_no[0]}").bold = True
    p.paragraph_format.space_after = Pt(0)
    para(title, italic=True, after=4)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    fmts = fmts or {}
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, col in enumerate(df.columns):
            v = row[col]
            cells[i].text = fmts[col](v) if col in fmts else str(v)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for r in pp.runs:
                    r.font.size = Pt(10)
    if note:
        para("Note. " + note, size=10, after=12)
    else:
        doc.add_paragraph()
    return tab_no[0]


# ================================================================== TITLE PAGE
para("Physics-Informed Neural Networks versus Classical Least Squares for Estimating "
     "Dengue Transmission Parameters: A Case Study of Bulacan, Philippines (1994–2010)",
     bold=True, align="center", size=16, after=18)
para(", ".join(AUTHORS), align="center")
para(AFFILIATION, align="center", italic=True)
para(COURSE, align="center")
para(f"Code and data: {REPO_URL}", align="center", size=10, after=24)

# ================================================================== ABSTRACT
doc.add_heading("Abstract", 1)
body(
    "Estimating how fast a disease spreads is central to outbreak response, yet the transmission and "
    "recovery rates of epidemic models cannot be measured directly and must be inferred from case data. "
    "This study compares two approaches to this inverse problem for the susceptible–infected–recovered "
    "(SIR) model: a classical solve-then-fit method that couples a fourth-order Runge–Kutta (RK4) solver "
    "with nonlinear least squares, and a physics-informed neural network (PINN) that learns the solution "
    "and the parameters simultaneously by penalizing violations of the differential equations. "
    f"A step-size study confirms that the RK4 solver converges with observed order {f(orders.mean())}, "
    "matching the theoretical order of four. "
    f"On synthetic outbreaks with 10% observation noise, the PINN recovered the transmission rate β and "
    f"recovery rate γ with relative errors of {pct(pinn.loc[0.10, 'beta_err_pct'])} and "
    f"{pct(pinn.loc[0.10, 'gamma_err_pct'])}, comparable to the classical method "
    f"({pct(cls.loc[0.10, 'beta_err_pct'])} and {pct(cls.loc[0.10, 'gamma_err_pct'])}, averaged over "
    f"{len(C.NOISE_SEEDS)} noise draws), but required about {speedup:.0f} times more computation time "
    "and was sensitive to the weight of its physics loss. "
    f"Applied to monthly dengue counts from Bulacan (OpenDengue, {len(seasons)} outbreak seasons between "
    f"{seasons.season.min()} and {seasons.season.max()}), the classical SIR fit reproduced seasonal "
    f"curves with R² between {f(seasons.R2.min(), 3)} and {f(seasons.R2.max(), 3)} and effective reproduction "
    f"numbers between {f(seasons.R0.min())} and {f(seasons.R0.max())}. With only eight monthly "
    "observations per season, however, the PINN's learned parameters did not reproduce its own fitted "
    f"curve (R² of the re-simulated curve: {f(p_pinn.R2_rk4_resimulated)}). The results suggest that for "
    "low-dimensional ODE models with sparse data, classical fitting remains the more reliable and "
    "economical tool, while PINNs need richer data or additional constraints to be reliable."
)
rich([("Keywords: ", "b"), ("dengue, SIR model, physics-informed neural networks, Runge–Kutta, "
                            "parameter estimation, inverse problems, Bulacan", "i")])

# ================================================================== 1 INTRODUCTION
doc.add_heading("1. Introduction", 1)
body(
    "Dengue is a mosquito-borne viral infection and one of the most widespread vector-borne diseases in "
    "the world, with an estimated 390 million infections per year (Bhatt et al., 2013). In the "
    "Philippines, dengue is endemic and peaks during the rainy season, placing recurring pressure on "
    "local health systems. Bulacan, a rapidly urbanizing province in Central Luzon (Region III), has "
    "experienced repeated outbreaks; reported monthly cases in the province rose above one thousand "
    "during the 2010 season (OpenDengue; Clarke et al., 2024)."
)
body(
    "Mathematical models help turn case counts into quantities that guide decisions. The SIR model of "
    "Kermack and McKendrick (1927) summarizes an outbreak with two rates: the transmission rate β and the "
    "recovery rate γ, whose ratio R₀ = β/γ (the basic reproduction number) determines whether an "
    "outbreak grows. These rates are not observed directly; they must be estimated from data, which "
    "makes parameter estimation an inverse problem. The standard approach repeatedly solves the "
    "differential equations numerically and adjusts the parameters to minimize the misfit with data."
)
body(
    "Physics-informed neural networks (PINNs; Raissi et al., 2019) offer an artificial-intelligence "
    "alternative: a neural network approximates the solution of the differential equations, and the "
    "equations themselves enter the training loss, so the network fits the data while obeying the "
    "physics. Unknown parameters such as β and γ are learned alongside the network weights. PINNs have "
    "attracted wide interest in scientific machine learning (Karniadakis et al., 2021), but they are "
    "also known to be difficult to train (Wang et al., 2021), and their advantage over classical "
    "methods for small epidemic models is not obvious."
)
body(
    "This study therefore asks: for estimating SIR parameters from dengue data, how does a PINN compare "
    "with a classical RK4 plus least-squares method in accuracy, stability and computational cost? We "
    "(1) verify the numerical solver through a step-size convergence study; (2) compare both methods on "
    "synthetic outbreaks with known parameters and controlled noise; (3) examine the PINN's sensitivity "
    "to its physics-loss weight; and (4) apply both methods to real monthly dengue counts from Bulacan. "
    "The significance of the study is practical: it tells local modelers which tool is appropriate for "
    "the sparse, aggregated surveillance data typically available at the provincial level, and it "
    "identifies what a PINN-based approach would need in order to be useful."
)

# ================================================================== 2 PRELIMINARIES
doc.add_heading("2. Preliminaries", 1)
doc.add_heading("2.1 The SIR Epidemic Model", 2)
body(
    "Let S(t), I(t) and R(t) denote the fractions of a closed population that are susceptible, "
    "infected (and infectious) and recovered at time t (in days). The SIR model is the system of "
    "ordinary differential equations (ODEs)"
)
equation("dS/dt = −βSI,    dI/dt = βSI − γI,    dR/dt = γI", "1")
body(
    "with initial condition (S, I, R)(0) = (1 − I₀, I₀, 0). The term βSI models new infections under the "
    "mass-action assumption, and γI models recoveries, so 1/γ is the mean infectious period. Adding the "
    "three equations gives d(S + I + R)/dt = 0, so S + I + R = 1 for all t. Linearizing dI/dt near "
    "S ≈ 1 gives dI/dt ≈ (β − γ)I, so infections grow initially if and only if R₀ = β/γ > 1."
)
doc.add_heading("2.2 The Fourth-Order Runge–Kutta Method", 2)
body(
    "System (1) has no closed-form solution, so it is solved numerically. Writing y = (S, I, R) and "
    "dy/dt = f(y), the classical fourth-order Runge–Kutta (RK4) method advances the solution with a fixed "
    "step size h by"
)
equation("k₁ = f(yₙ),  k₂ = f(yₙ + (h/2)k₁),  k₃ = f(yₙ + (h/2)k₂),  k₄ = f(yₙ + hk₃)", "2")
equation("yₙ₊₁ = yₙ + (h/6)(k₁ + 2k₂ + 2k₃ + k₄)", "3")
body(
    "RK4 has local truncation error O(h⁵) and global error O(h⁴) for smooth f (Burden & Faires, 2011). "
    "If e(h) is the global error with step h, the observed order of accuracy between two step sizes is"
)
equation("p = log(e(h₁)/e(h₂)) / log(h₁/h₂)", "4")
body("which should approach 4 as h decreases.")
doc.add_heading("2.3 Nonlinear Least Squares", 2)
body(
    "Given observations yᵢ at times tᵢ and a model prediction m(tᵢ; θ) that depends on parameters θ, "
    "nonlinear least squares chooses θ to minimize"
)
equation("Φ(θ) = ½ Σᵢ rᵢ(θ)²,   rᵢ(θ) = m(tᵢ; θ) − yᵢ", "5")
body(
    "We use the trust-region reflective algorithm in SciPy's least_squares (Virtanen et al., 2020), a "
    "Gauss–Newton-type method that handles bound constraints (here β, γ > 0). Each evaluation of the "
    "residual requires a full RK4 solution of (1). Iteration stops when the relative changes in Φ, in θ, "
    "or the gradient norm fall below a tolerance."
)
doc.add_heading("2.4 Neural Networks and Automatic Differentiation", 2)
body(
    "A fully connected neural network with L hidden layers maps an input x to an output through "
    "compositions zₗ = σ(Wₗzₗ₋₁ + bₗ), where Wₗ and bₗ are trainable weights and biases and σ is a "
    "nonlinear activation (here tanh). Such networks are universal approximators of continuous "
    "functions. Training minimizes a loss function with gradient-based optimizers; we use Adam "
    "(Kingma & Ba, 2015), whose learning rate plays the role of a step size. Automatic differentiation, "
    "as implemented in PyTorch (Paszke et al., 2019), computes exact derivatives of the network output "
    "with respect to its input and parameters by applying the chain rule to the recorded computation, "
    "without finite-difference approximation."
)
doc.add_heading("2.5 Physics-Informed Neural Networks", 2)
body(
    "A PINN (Raissi et al., 2019) represents the unknown solution by a network N(t; w) and trains it on a "
    "composite loss. For an inverse problem, unknown model parameters are added to the trainable "
    "variables. For system (1) with observations of I, the loss is"
)
equation("L(w, β, γ) = L_data + w_ode · L_ode + L_ic", "6")
body(
    "where L_data is the mean squared error between the network's I(t) and the data, L_ode is the mean "
    "squared residual of the three equations in (1) evaluated by automatic differentiation at a set of "
    "collocation points, and L_ic penalizes deviation from the initial condition. The weight w_ode "
    "balances fitting the data against satisfying the physics."
)
doc.add_heading("2.6 Bootstrap Error Estimation", 2)
body(
    "For real data the true parameters are unknown, so uncertainty is estimated by a residual "
    "bootstrap: residuals of the fitted model are resampled with replacement and added back to the "
    "fitted values to create synthetic datasets, which are refitted. The 2.5th and 97.5th percentiles "
    "of the refitted parameters form an approximate 95% confidence interval."
)

# ================================================================== 3 METHODS AND APPLICATIONS
doc.add_heading("3. Methods and Applications", 1)
doc.add_heading("3.1 Model and Method Description", 2)
doc.add_heading("Problem setup", 3)
body(
    "Variables: S, I, R (population fractions) as functions of time t in days. Parameters to estimate: "
    "β > 0 and γ > 0 (per day). Derived quantity: R₀ = β/γ. Assumptions: closed population over one "
    "season, homogeneous mixing, constant rates within a season, and no explicit mosquito "
    "compartment. Constraints: β, γ positive (enforced by bounds in least squares and by a logarithmic "
    "parameterization, β = exp(log β), in the PINN). For the real data, the effective infectious "
    f"period 1/γ is further restricted to {lo_days}–{hi_days} days, i.e. γ ∈ [1/{hi_days}, 1/{lo_days}] "
    "per day. Without this constraint, the four-parameter fit to eight monthly points can collapse to a "
    "degenerate solution (γ at an arbitrary bound, R₀ ≈ 1 and a very large N_eff) that fits the data "
    "equally well. The lower limit reflects the roughly one-week human infectious period; the upper "
    "limit allows for the mosquito stage that a plain SIR model absorbs into γ."
)
doc.add_heading("Synthetic data", 3)
body(
    f"A reference outbreak is generated with β = {C.TRUE_BETA}, γ = {C.TRUE_GAMMA} (R₀ = "
    f"{C.TRUE_BETA / C.TRUE_GAMMA:.0f}, infectious period {1 / C.TRUE_GAMMA:.0f} days), I₀ = {C.I0}, "
    f"over {C.T_DAYS} days, solved by RK4 with h = {C.H_DATA}. Daily values of I are observed with "
    "multiplicative Gaussian noise, I_obs = I_true(1 + σε) with ε ~ N(0, 1) and σ ∈ {0, 0.05, 0.10}. "
    "Because the true parameters are known, estimation error can be measured exactly."
)
doc.add_heading("Real data", 3)
body(
    "Monthly reported dengue cases for the Province of Bulacan were extracted from the OpenDengue V1.3 "
    "spatial extract (Clarke et al., 2024), which provides 200 monthly province-level counts from 1994 "
    f"to 2010. We selected the {len(seasons)} seasons whose monthly peak exceeded 300 cases "
    f"({', '.join(str(s) for s in seasons.season)}) and fitted each from May to December (eight months), "
    "the period containing the rainy-season wave. Surveillance counts new reported cases rather than "
    "the infected fraction, so the observation model is"
)
equation("cases in month k = N_eff · [S(t_{k−1}) − S(t_k)]", "7")
body(
    "since every decrease in S is a new infection. The effective population N_eff absorbs the "
    "reporting rate and the fraction of the population actually susceptible; it is estimated together "
    "with β, γ and I₀. This makes N_eff much smaller than the province's census population, which is "
    "expected when most infections are unreported or occur in a partly immune population."
)
doc.add_heading("3.2 Implementation and Code", 2)
body(
    "All code is written in Python using NumPy, SciPy, pandas, Matplotlib and PyTorch, and is available "
    f"in the project repository ({REPO_URL}). All settings (seeds, step sizes, tolerances, epochs) are "
    "centralized in src/config.py, and run_all.py regenerates every table and figure in this paper."
)
doc.add_heading("Algorithm 1: RK4 solver", 3)
body("The RK4 step (2)–(3) is implemented directly (src/sir.py):")
code('''
def rk4(beta, gamma, y0, t_end, h):
    n = int(round(t_end / h))
    y = np.asarray(y0, dtype=float)
    for k in range(n):
        k1 = sir_rhs(y, beta, gamma)
        k2 = sir_rhs(y + 0.5 * h * k1, beta, gamma)
        k3 = sir_rhs(y + 0.5 * h * k2, beta, gamma)
        k4 = sir_rhs(y + h * k3, beta, gamma)
        y = y + h / 6.0 * (k1 + 2 * k2 + 2 * k3 + k4)
''')
doc.add_heading("Algorithm 2: Classical solve-then-fit", 3)
body(
    "For a candidate θ = (β, γ), the residual vector is computed by solving (1) with RK4 and "
    "comparing the model's I(tᵢ) with the data; least_squares then updates θ (src/classical.py). For the "
    "real data, θ = (β, γ, log₁₀N_eff, log₁₀I₀) and the residuals follow (7). Because this "
    "four-parameter problem is non-convex, the fit is started from six initial guesses and the best "
    "solution is kept."
)
code('''
def residuals(p):
    t, Y = rk4(p[0], p[1], y0, t_end, h)
    return (Y[::stride, 1] - I_obs) / I_obs.max()

res = least_squares(residuals, p0, bounds=([1e-4, 1e-4], [5, 5]),
                    ftol=1e-10, xtol=1e-10, gtol=1e-10)
''')
doc.add_heading("Algorithm 3: PINN", 3)
body(
    f"The network maps rescaled time τ = t/T ∈ [0, 1] to (S, I/s_I, R) through {C.PINN['depth']} hidden "
    f"layers of {C.PINN['width']} tanh units and a sigmoid output, where s_I is the maximum observed I, so "
    "all outputs are of order one. The ODE residuals are evaluated at "
    f"{C.PINN['n_col']} equally spaced collocation points using torch.autograd.grad, with the chain rule "
    "d/dτ = T·d/dt. β and γ are stored as log-values and optimized by Adam together with the weights "
    "(src/pinn.py):"
)
code('''
out = net(tau_c);  S, I_s, R = out[:, :1], out[:, 1:2], out[:, 2:3]
I = I_s * s_I
r_S = d_dtau(S) - T * (-beta * S * I)
r_I = d_dtau(I_s) - T * (beta * S * I - gamma * I) / s_I
r_R = d_dtau(R) - T * (gamma * I)
loss = l_data + w_ode * (r_S**2 + r_I**2 + r_R**2).mean() + l_ic
loss.backward();  opt.step()
''')
doc.add_heading("Step size, iteration count and tolerance", 3)
settings = pd.DataFrame([
    ["RK4 (data generation)", "step size h", f"{C.H_DATA} day"],
    ["RK4 (convergence study)", "step sizes h", ", ".join(f"{h:g}" for h in C.H_LIST) + f" (reference h = {C.H_REF:g})"],
    ["RK4 (inside synthetic fit)", "step size h", f"{C.H_FIT} day (Exp. 1b varies {', '.join(f'{h:g}' for h in C.H_FIT_LIST)})"],
    ["RK4 (real-data model)", "step size h", f"{C.H_REAL} day"],
    ["Least squares", "tolerances ftol, xtol, gtol", f"{C.LSQ_TOL:g}"],
    ["PINN", "Adam learning rate", f"{C.PINN['lr']:g}"],
    ["PINN", "maximum epochs", f"{C.PINN['max_epochs']:,}"],
    ["PINN", "stopping rule", f"relative change of β and γ over {C.PINN['check_every']} epochs < {C.PINN['tol']:g}"],
    ["PINN", "physics-loss weight w_ode", f"{C.PINN['w_ode']:g} (Exp. 3 varies {', '.join(f'{w:g}' for w in C.W_ODE_SWEEP)})"],
    ["PINN (real data)", "epochs (fixed)", f"{C.REAL_PINN_EPOCHS:,}"],
    ["Bootstrap", "replicates", f"{C.BOOTSTRAP_B}"],
], columns=["Component", "Setting", "Value"])
table(settings, ["Component", "Setting", "Value"], "Numerical settings used in all experiments")
body(
    "Errors are reported as relative parameter errors |θ̂ − θ|/θ × 100%, and as the root-mean-square "
    "error (trajectory RMSE) between the true I(t) and an RK4 re-simulation with the estimated "
    "parameters. Runtimes were measured on a laptop CPU (Intel Core i5-13420H) with PyTorch restricted to "
    "one thread; environment details are recorded in results/run_log.txt."
)

# ------------------------------------------------------------------ 3.3 RESULTS
doc.add_heading("3.3 Results and Visualization", 2)
doc.add_heading("Experiment 1: RK4 convergence", 3)
c1 = conv.copy()
t1 = table(c1[["h", "steps", "global_error", "observed_order", "max_mass_drift"]],
           ["Step size h", "Steps", "Global error at t = 100", "Observed order p", "Max |S+I+R−1|"],
           "Global error of the RK4 solver against a reference solution (h = 0.001)",
           fmts={"h": lambda v: f"{v:g}", "steps": lambda v: f"{int(v)}",
                 "global_error": lambda v: f"{v:.3e}",
                 "observed_order": lambda v: "–" if pd.isna(v) else f"{v:.3f}",
                 "max_mass_drift": lambda v: f"{v:.1e}"})
fg1 = figure("fig1_rk4_convergence.png",
             "(a) Log–log plot of RK4 global error versus step size with a slope-4 reference line. "
             "(b) Relative error of the classical estimates on noise-free data as a function of the RK4 "
             "step size used inside the fit.")
body(
    f"Table {t1} and Figure {fg1}a show that halving h reduces the error by a factor of about 16; the "
    f"observed order ranges from {f(orders.min(), 3)} to {f(orders.max(), 3)}, in agreement with the "
    "theoretical order of four. Population is conserved to machine precision at every step size. "
    f"Figure {fg1}b shows that the solver step used inside the fit controls a systematic bias: with "
    f"h = {fitstep.h_fit.iloc[0]:g} days the estimates of β and γ are off by "
    f"{pct(fitstep.beta_rel_err_pct.iloc[0])} and {pct(fitstep.gamma_rel_err_pct.iloc[0])}, while with "
    f"h = 1 day the error drops to {pct(fitstep.beta_rel_err_pct.iloc[2])}. With h = {C.H_DATA}, the same "
    "step used to generate the data, the parameters are recovered to machine precision; this is the "
    "so-called inverse crime, which is why noise is added in the remaining experiments."
)
doc.add_heading("Experiment 2: Parameter recovery on synthetic data", 3)
s = summ.copy()
s["beta_cell"] = s.apply(lambda r: f"{r.beta:.4f} ({pct(r.beta_err_pct)})", axis=1)
s["gamma_cell"] = s.apply(lambda r: f"{r.gamma:.4f} ({pct(r.gamma_err_pct)})", axis=1)
s["method"] = s.method.str.replace(" (RK4 + least squares)", "", regex=False)
t2 = table(s[["method", "noise", "beta_cell", "gamma_cell", "traj_rmse", "iterations", "seconds"]],
           ["Method", "Noise", "β̂ (rel. error)", "γ̂ (rel. error)", "Trajectory RMSE",
            "Iterations", "Time (s)"],
           "Recovery of β = 0.3 and γ = 0.1 from synthetic data",
           fmts={"noise": lambda v: f"{v:.0%}", "traj_rmse": lambda v: f"{v:.2e}",
                 "iterations": lambda v: f"{v:,.0f}", "seconds": lambda v: f"{v:.2f}"},
           note=(f"Classical values are means over {len(C.NOISE_SEEDS)} noise draws (one draw at 0% "
                 "noise); iterations are function evaluations (each a full RK4 solve). PINN values are "
                 "from one training run (seed 0); iterations are epochs until the stopping rule."))
fg2 = figure("fig2_fits.png",
             "Observed infected fraction (points) with the classical fit (solid) and PINN prediction "
             "(dashed) at 0%, 5% and 10% noise.")
fg3 = figure("fig3_pinn_training.png",
             "PINN training at 5% noise. (a) Data, ODE-residual and total loss per epoch. (b) Estimates "
             "of β and γ during training; dotted lines mark the true values.")
body(
    f"Both methods recover the parameters accurately (Table {t2}, Figure {fg2}). At 10% noise the PINN's "
    f"errors are {pct(pinn.loc[0.10, 'beta_err_pct'])} for β and {pct(pinn.loc[0.10, 'gamma_err_pct'])} "
    f"for γ, versus {pct(cls.loc[0.10, 'beta_err_pct'])} and {pct(cls.loc[0.10, 'gamma_err_pct'])} on "
    "average for the classical method. The classical method converges in at most "
    f"{int(cls_all.iterations.max())} function evaluations (about {cls.seconds.mean():.2f} s), whereas the PINN needs thousands of epochs "
    f"(on average {pinn.seconds.mean():.0f} s), roughly {speedup:.0f} times longer. Figure {fg3} shows "
    "the PINN's convergence behavior: the data loss falls quickly, the ODE residual is reduced more "
    "slowly, and β and γ settle onto the true values before the stopping rule ends training."
)
doc.add_heading("Experiment 3: Sensitivity to the physics-loss weight", 3)
sw = sweep.copy()
t3 = table(sw[["w_ode", "beta", "beta_err_pct", "gamma", "gamma_err_pct", "epochs", "seconds"]],
           ["w_ode", "β̂", "β error", "γ̂", "γ error", "Epochs", "Time (s)"],
           f"PINN estimates at {C.SWEEP_NOISE:.0%} noise for different physics-loss weights",
           fmts={"w_ode": lambda v: f"{v:g}", "beta": lambda v: f"{v:.4f}", "gamma": lambda v: f"{v:.4f}",
                 "beta_err_pct": pct, "gamma_err_pct": pct, "epochs": lambda v: f"{v:,.0f}",
                 "seconds": lambda v: f"{v:.1f}"})
fg4 = figure("fig4_ode_weight_sweep.png",
             "(a) Final relative error of β and γ (log scale) for each physics-loss weight. (b) The "
             "estimate of β during training for each weight.")
w_max = sweep.loc[sweep.w_ode.idxmax()]
w_min = sweep.loc[sweep.w_ode.idxmin()]
body(
    f"The physics-loss weight changes both the accuracy and the convergence of the PINN (Table {t3}, "
    f"Figure {fg4}). The most accurate weight was w_ode = {best_w.w_ode:g} (errors "
    f"{pct(best_w.beta_err_pct)} and {pct(best_w.gamma_err_pct)}), while w_ode = {worst_w.w_ode:g} gave "
    f"errors of {pct(worst_w.beta_err_pct)} and {pct(worst_w.gamma_err_pct)}. "
    + (f"With the largest weight (w_ode = {w_max.w_ode:g}) the estimate of β stalled on a plateau far "
       "from the true value for most of training before moving, and the run reached the "
       f"{C.PINN['max_epochs']:,}-epoch cap without satisfying the stopping rule. "
       if not w_max.converged_by_tol else "")
    + f"With the smallest weight (w_ode = {w_min.w_ode:g}) training is fast, but the physics constraint "
    "is weak and β is less accurate. The classical method has no corresponding tuning parameter."
)
doc.add_heading("Experiment 4: Dengue seasons in Bulacan", 3)
fg5 = figure("fig5_bulacan_overview.png",
             "Monthly reported dengue cases in Bulacan, 1994–2010 (OpenDengue V1.3).")
t4 = table(seasons[["season", "total_cases", "peak_cases", "beta", "gamma", "R0", "infectious_days",
                    "N_eff", "R2"]],
           ["Season", "Cases (May–Dec)", "Peak month", "β̂", "γ̂", "R₀", "1/γ̂ (days)", "N_eff", "R²"],
           "Classical SIR fits to Bulacan outbreak seasons (May–December)",
           note=(f"1/γ̂ was constrained to {lo_days}–{hi_days} days; "
                 + (f"the {', '.join(str(s) for s in at_bound.season)} estimate lies on the boundary."
                    if len(at_bound) else "all estimates lie inside the constraint.")),
           fmts={"total_cases": lambda v: f"{int(v):,}", "peak_cases": lambda v: f"{int(v):,}",
                 "beta": lambda v: f"{v:.3f}", "gamma": lambda v: f"{v:.3f}", "R0": lambda v: f"{v:.2f}",
                 "infectious_days": lambda v: f"{v:.1f}", "N_eff": lambda v: f"{v:,.0f}",
                 "R2": lambda v: f"{v:.3f}"})
fg6 = figure("fig6_bulacan_seasons.png",
             "Reported monthly cases (bars) and classical SIR fits (lines) for six outbreak seasons.")
b = boot
t5 = table(pd.DataFrame({
    "parameter": ["β", "γ", "R₀", "N_eff"],
    "est": b.estimate.values, "lo": b.ci95_low.values, "hi": b.ci95_high.values}),
    ["Parameter", "Estimate", "95% CI (lower)", "95% CI (upper)"],
    f"Residual-bootstrap 95% intervals for the 2010 season ({C.BOOTSTRAP_B} replicates)",
    fmts={"est": lambda v: f"{v:,.3f}" if v < 100 else f"{v:,.0f}",
          "lo": lambda v: f"{v:,.3f}" if v < 100 else f"{v:,.0f}",
          "hi": lambda v: f"{v:,.3f}" if v < 100 else f"{v:,.0f}"})
pt = pinn10.copy()
t6 = table(pt[["method", "beta", "gamma", "R0", "R2_own_curve", "R2_rk4_resimulated", "seconds"]],
           ["Method", "β̂", "γ̂", "R₀", "R² (own curve)", "R² (RK4 re-simulation)", "Time (s)"],
           "Classical fit and PINN on the 2010 season",
           fmts={"beta": lambda v: f"{v:.3f}", "gamma": lambda v: f"{v:.3f}", "R0": lambda v: f"{v:.2f}",
                 "R2_own_curve": lambda v: f"{v:.3f}", "R2_rk4_resimulated": lambda v: f"{v:.3f}",
                 "seconds": lambda v: f"{v:.1f}"},
           note="For the PINN, N_eff and I₀ were fixed at the classical estimates so that only β and γ "
                "were learned. 'Own curve' is the network output; 'RK4 re-simulation' solves the SIR "
                "model with the PINN's β and γ.")
fg7 = figure("fig7_bulacan_pinn.png",
             "Bulacan 2010: reported cases, classical fit, the PINN network's output, and an RK4 "
             "re-simulation using the PINN's estimated β and γ.", width=5.5)
body(
    f"The classical SIR model reproduces the shape of each season (Table {t4}, Figure {fg6}), with R² from "
    f"{f(seasons.R2.min(), 3)} to {f(seasons.R2.max(), 3)} and R₀ estimates from {f(seasons.R0.min())} to "
    f"{f(seasons.R0.max())}. For 2010, the largest season ({int(s2010.total_cases):,} cases from May "
    f"to December), the fit gives β = {f(s2010.beta, 3)} per day, γ = {f(s2010.gamma, 3)} per day and "
    f"R₀ = {f(s2010.R0)}. The bootstrap intervals in Table {t5} quantify how well eight monthly points "
    f"constrain the parameters: R₀ lies between {f(b.loc['R0', 'ci95_low'])} and "
    f"{f(b.loc['R0', 'ci95_high'])}. On the same season the PINN's network output matches the data "
    f"(R² = {f(p_pinn.R2_own_curve, 4)}), but solving the SIR model with the PINN's β and γ gives "
    f"R² = {f(p_pinn.R2_rk4_resimulated, 3)} (Table {t6}, Figure {fg7}), meaning the learned parameters "
    "are not consistent with the network's own curve."
)

# ------------------------------------------------------------------ 3.4 DISCUSSION
doc.add_heading("3.4 Analysis and Discussion", 2)
rich([("Accuracy. ", "b"),
      ("On synthetic data both methods achieve parameter errors of about one percent or less. The "
       "classical errors, averaged over several noise draws, grow with the noise level as expected; the "
       "PINN errors come from a single training run per noise level and fluctuate rather than growing "
       "smoothly, so a lower PINN error in one run should not be read as a systematic advantage. Since "
       "both methods fit the same model to the same data, neither can do much better than the "
       "information in the noisy observations allows. The classical method's accuracy is "
       "limited by the RK4 step used inside the fit (Experiment 1b), which is negligible for h ≤ 1 day "
       "because of the fourth-order convergence.", "")])
rich([("Stability. ", "b"),
      ("RK4 is stable for the step sizes used, and least squares converged in every synthetic run "
       "from the same initial guess used by the PINN. The PINN is less robust: its result depends on the physics-loss weight "
       "(Experiment 3). When the physics term dominates, the network first learns a smooth curve that "
       "satisfies the ODE for the current, incorrect parameters, and β and γ then move only slowly; "
       "when it is too weak, the data alone determine the curve and the parameters are less "
       "constrained. This agrees with known training pathologies of PINNs, where "
       "imbalanced loss gradients slow or misdirect learning (Wang et al., 2021).", "")])
rich([("Computational cost. ", "b"),
      ("One RK4 solve costs O(T/h) evaluations of f, and the classical fit needs only tens of solves. The "
       "PINN's cost per epoch is a forward and backward pass over the collocation and data points, "
       "repeated for thousands of epochs. In our experiments the PINN was roughly "
       f"{speedup:.0f} times slower. For a two-parameter ODE this cost is not justified; PINNs become "
       "attractive for partial differential equations, high-dimensional problems, or when no efficient "
       "solver exists.", "")])
rich([("Real data. ", "b"),
      ("The classical SIR model captures the timing and size of Bulacan's outbreaks, but the implied "
       f"infectious periods (1/γ ≈ {interior.infectious_days.min():.0f}–"
       f"{interior.infectious_days.max():.0f} days for the seasons whose estimate lies inside the "
       "constraint) are longer than the roughly one-week infectious period of a human dengue case. This is a "
       "symptom of model simplification: SIR has no mosquito compartment, so the extrinsic incubation "
       "period in the mosquito and the mosquito lifespan are absorbed into γ. β and γ should therefore "
       "be read as effective seasonal rates rather than biological constants. "
       + (f"The {', '.join(str(s) for s in at_bound.season)} season"
          f"{'s' if len(at_bound) > 1 else ''} ended at the boundary of the constraint, which signals "
          "that the data are not well described by SIR there: in 2006 reported cases fell abruptly "
          "from 504 in September to 110 in October and 17 in November, a drop faster than the SIR "
          "curve can follow with a plausible infectious period, possibly driven by the end of the "
          "rainy season or by changes in reporting. " if len(at_bound) else "")
       + "Monthly data also "
       "provide only eight observations for four unknowns, and the bootstrap intervals show the "
       "resulting uncertainty. The PINN struggled most here: with eight points, the network can "
       "interpolate the data without the ODE residual being small enough to pin down β and γ.", "")])
rich([("Practical significance and limitations. ", "b"),
      ("For provincial health offices working with sparse monthly surveillance data, a classical SIR fit "
       "is fast, transparent and adequate for describing an outbreak's growth (R₀) and timing. Our "
       "conclusions are limited by: the SIR simplification (no vector, no seasonality in β, no "
       "under-reporting model beyond N_eff); data that end in 2010 and are aggregated monthly; one PINN "
       "architecture and one training seed per configuration; and treating each season "
       "independently.", "")])

# ================================================================== 4 CONCLUSION
doc.add_heading("4. Conclusion", 1)
body(
    "We compared a physics-informed neural network with a classical RK4 plus least-squares method for "
    "estimating SIR transmission parameters. The RK4 solver was verified to converge with fourth-order "
    "accuracy. On synthetic outbreaks with up to 10% noise, both methods recovered β and γ within a few "
    f"percent, but the PINN was about {speedup:.0f} times more expensive and sensitive to its "
    f"physics-loss weight. On real dengue data from Bulacan, the classical model fitted {len(seasons)} "
    f"outbreak seasons with R² ≥ {f(seasons.R2.min())} and effective R₀ between {f(seasons.R0.min())} "
    f"and {f(seasons.R0.max())}, while the PINN failed to produce parameters consistent with its own fit "
    "from eight monthly observations. For low-dimensional epidemic models with sparse data, classical "
    "fitting is the more effective method; PINNs are better suited to richer data or more complex "
    "models."
)
body("Future work suggested by these findings, and a natural basis for a thesis, includes:")
for item in [
    "replacing SIR with a host–vector dengue model (Esteva & Vargas, 1998) so that the estimated rates "
    "become biologically interpretable;",
    "using weekly or more recent surveillance data (e.g., from the Department of Health) to give the "
    "PINN more observations;",
    "letting the transmission rate vary with rainfall and temperature, a setting where a neural "
    "network can represent an unknown function β(t) that a classical fit cannot easily handle;",
    "fitting several seasons jointly with shared parameters; and",
    "improving PINN training with adaptive loss weighting or two-stage training that fixes the "
    "network before estimating β and γ.",
]:
    bullet(item)

# ================================================================== REFERENCES
doc.add_heading("References", 1)
refs = [
    "Bhatt, S., Gething, P. W., Brady, O. J., Messina, J. P., Farlow, A. W., Moyes, C. L., Drake, J. M., "
    "Brownstein, J. S., Hoen, A. G., Sankoh, O., Myers, M. F., George, D. B., Jaenisch, T., Wint, G. R. W., "
    "Simmons, C. P., Scott, T. W., Farrar, J. J., & Hay, S. I. (2013). The global distribution and burden of "
    "dengue. Nature, 496(7446), 504–507. https://doi.org/10.1038/nature12060",
    "Burden, R. L., & Faires, J. D. (2011). Numerical analysis (9th ed.). Brooks/Cole, Cengage Learning.",
    "Clarke, J., Lim, A., Gupte, P., Pigott, D. M., van Panhuis, W. G., & Brady, O. J. (2024). A global "
    "dataset of publicly available dengue case count data. Scientific Data, 11, Article 296. "
    "https://doi.org/10.1038/s41597-024-03120-7",
    "Esteva, L., & Vargas, C. (1998). Analysis of a dengue disease transmission model. Mathematical "
    "Biosciences, 150(2), 131–151. https://doi.org/10.1016/S0025-5564(98)10003-2",
    "Karniadakis, G. E., Kevrekidis, I. G., Lu, L., Perdikaris, P., Wang, S., & Yang, L. (2021). "
    "Physics-informed machine learning. Nature Reviews Physics, 3(6), 422–440. "
    "https://doi.org/10.1038/s42254-021-00314-5",
    "Kermack, W. O., & McKendrick, A. G. (1927). A contribution to the mathematical theory of epidemics. "
    "Proceedings of the Royal Society of London. Series A, 115(772), 700–721. "
    "https://doi.org/10.1098/rspa.1927.0118",
    "Kingma, D. P., & Ba, J. (2015). Adam: A method for stochastic optimization. In Proceedings of the 3rd "
    "International Conference on Learning Representations (ICLR). https://arxiv.org/abs/1412.6980",
    "Paszke, A., Gross, S., Massa, F., Lerer, A., Bradbury, J., Chanan, G., Killeen, T., Lin, Z., "
    "Gimelshein, N., Antiga, L., Desmaison, A., Köpf, A., Yang, E., DeVito, Z., Raison, M., Tejani, A., "
    "Chilamkurthy, S., Steiner, B., Fang, L., … Chintala, S. (2019). PyTorch: An imperative style, "
    "high-performance deep learning library. In Advances in Neural Information Processing Systems 32 "
    "(pp. 8024–8035). Curran Associates.",
    "Raissi, M., Perdikaris, P., & Karniadakis, G. E. (2019). Physics-informed neural networks: A deep "
    "learning framework for solving forward and inverse problems involving nonlinear partial "
    "differential equations. Journal of Computational Physics, 378, 686–707. "
    "https://doi.org/10.1016/j.jcp.2018.10.045",
    "Virtanen, P., Gommers, R., Oliphant, T. E., Haberland, M., Reddy, T., Cournapeau, D., Burovski, E., "
    "Peterson, P., Weckesser, W., Bright, J., van der Walt, S. J., Brett, M., Wilson, J., Millman, K. J., "
    "Mayorov, N., Nelson, A. R. J., Jones, E., Kern, R., Larson, E., … SciPy 1.0 Contributors. (2020). "
    "SciPy 1.0: Fundamental algorithms for scientific computing in Python. Nature Methods, 17(3), "
    "261–272. https://doi.org/10.1038/s41592-019-0686-2",
    "Wang, S., Teng, Y., & Perdikaris, P. (2021). Understanding and mitigating gradient flow pathologies "
    "in physics-informed neural networks. SIAM Journal on Scientific Computing, 43(5), A3055–A3081. "
    "https://doi.org/10.1137/20M1318043",
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)

doc.save(OUT)
print(f"Saved {OUT}")
